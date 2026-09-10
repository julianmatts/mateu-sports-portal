# -*- coding: utf-8 -*-
# Genera regalias/siembra-2026.js con el histórico oficial de la temporada:
#   - meses: unidades/regalía/importe por mes, tomados a mano de los CONSOLIDADO
#     de "Regalias Ruge SE y EDLP <mes> 2026.xlsx" (dict `meses` de abajo).
#   - rda: detalle mensual de Ruge Deportes Amateur, leído de los Excel
#     "DETALLE VENTA RUGE DEPORTES AMATEURS..." (rutas RDA_MAY/RDA_JUN).
# Correr con Python 3 + openpyxl en la máquina de Juli:  python scripts/gen-siembra-regalias.py
# Si se agrega un mes: sumar la entrada a `meses`, apuntar el RDA nuevo y subir
# SIEMBRA_V en regalias/index.html para que los navegadores re-siembren.
import openpyxl, json, unicodedata, os

RDA_JUN = r"C:\Users\julia\OneDrive\Desktop\EDLP 2026\Regalias ruge\Regalias Junio 2026\DETALLE VENTA RUGE DEPORTES AMATEURS CANAL MATEU JUNIO 2026.xlsx"
RDA_MAY = r"C:\Users\julia\OneDrive\Desktop\EDLP 2026\Regalias ruge\Regalias Mayo 2026\DETALLE VENTA RUGE DEPORTES AMATEURS CANAL MATEU MAYO 2026.xlsx"
MES_NUM = {'ENERO':1,'FEBRERO':2,'MARZO':3,'ABRIL':4,'MAYO':5,'JUNIO':6,'JULIO':7,'AGOSTO':8,'SEPTIEMBRE':9,'OCTUBRE':10,'NOVIEMBRE':11,'DICIEMBRE':12}

def leer_rda(p):
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['DETALLE MES']
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = None; hi = 0
    for i,r in enumerate(rows[:10]):
        up=[str(x).upper().strip() if x else '' for x in r]
        if 'MES' in up and any('IMPORTE' in c for c in up): hdr=up; hi=i; break
    print(os.path.basename(p), '-> header:', hdr)
    idx = {}
    for j,c in enumerate(hdr):
        cn = unicodedata.normalize('NFD', c).encode('ascii','ignore').decode()
        if cn=='MES': idx['mes']=j
        elif 'RUBRO' in cn: idx['rubro']=j
        elif 'GRUPO' in cn: idx['g1']=j
        elif 'LISTA' in cn: idx['lista']=j
        elif 'ARTICULO' in cn or 'DESCRIP' in cn: idx['art']=j
        elif 'BARRAS' in cn or 'CODIGO' in cn or cn=='COD': idx['cod']=j
        elif 'CANTIDAD' in cn or cn=='CANT': idx['cant']=j
        elif 'IMPORTE' in cn or 'MONTO' in cn: idx['imp']=j
    print('   idx:', idx)
    acc = {}
    for r in rows[hi+1:]:
        if r is None or all(v is None or str(v).strip()=='' for v in r): continue
        mesv = str(r[idx['mes']] or '').upper().strip()
        mm = MES_NUM.get(mesv)
        art = str(r[idx['art']] or '').strip()
        if not mm or not art: continue
        key = "2026-%02d" % mm
        cant = round(float(r[idx['cant']] or 0)); imp = float(r[idx['imp']] or 0)
        a = acc.setdefault(key, {'unidades':0,'importe':0.0,'lineas':[]})
        a['unidades'] += cant; a['importe'] += imp
        a['lineas'].append({'rubro':str(r[idx['rubro']] or '') if 'rubro' in idx else '',
                            'g1':str(r[idx['g1']] or '') if 'g1' in idx else '',
                            'lista':str(r[idx['lista']] or '') if 'lista' in idx else '',
                            'articulo':art,
                            'cod':str(r[idx['cod']] or '') if 'cod' in idx else '',
                            'cant':cant, 'importe':round(imp,2)})
    for k,a in acc.items():
        a['importe'] = round(a['importe'],2)
        a['regalia'] = round(a['importe']*0.05,2)
    return acc

rda = leer_rda(RDA_MAY)          # ene-may (regalias 5% ok)
rda_jun = leer_rda(RDA_JUN)
for k,v in rda_jun.items():
    if k not in rda or k=='2026-06': rda[k]=v
print('RDA meses:', {k:(v['unidades'],v['importe'],v['regalia'],len(v['lineas'])) for k,v in sorted(rda.items())})

meses = {
 '2026-01': {'edlpUnidades':7954,'rugeUnidades':393,'retail':6469,'mayor':4208,'unidades':10677,'regalia':73929600.92,'importe':504693132.88},
 '2026-02': {'edlpUnidades':4537,'rugeUnidades':614,'retail':6053,'mayor':799, 'unidades':6852, 'regalia':36336238.93,'importe':341733577.91},
 '2026-03': {'edlpUnidades':4618,'rugeUnidades':557,'retail':5142,'mayor':1952,'unidades':7094, 'regalia':43982372.01,'importe':351571604.03},
 '2026-04': {'edlpUnidades':6351,'rugeUnidades':307,'retail':5414,'mayor':2498,'unidades':7912, 'regalia':84784137.15,'importe':514199682.30},
 '2026-05': {'edlpUnidades':4375,'rugeUnidades':630,'retail':5959,'mayor':658, 'unidades':6617, 'regalia':82369718.66,'importe':468401719.27},
 '2026-06': {'edlpUnidades':4633,'rugeUnidades':562,'retail':5857,'mayor':799, 'unidades':6656, 'regalia':74889584.80,'importe':417287788.34},
 '2026-07': {'edlpUnidades':2420,'rugeUnidades':409,'retail':3651,'mayor':166, 'unidades':3817, 'regalia':41897173.81,'importe':231485661.97},
}

# Historial de unidades por concepto = cuadro A:H del CONSOLIDADO. Sale del último
# Excel de liquidación a mano, que ya trae todos los meses anteriores. El Excel que
# exporta el módulo lo repite en valores (los meses que se liquidan en la app se
# guardan en el ledger con el mismo formato, campo `conc`).
HIST = r"C:\Users\julia\OneDrive\Desktop\EDLP 2026\Regalias ruge\Regalias Julio 2026\Regalias Ruge SE y EDLP Julio 2026.xlsx"

def leer_consolidado(p):
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['CONSOLIDADO']
    n = lambda v: int(round(float(v))) if v not in (None, '') else 0
    out = {}
    for r in ws.iter_rows(min_row=2, max_row=25, max_col=8, values_only=True):
        mm = MES_NUM.get(str(r[0] or '').upper().strip())
        concepto = str(r[1] or '').upper().strip()
        if not mm or all(v in (None, '') for v in r[2:8]): continue
        d = out.setdefault("2026-%02d" % mm, {'may':0,'edlp':0,'ruge':0,'acc':0,'ruge2da':0,'edlp2da':0,'edlpAnt':0})
        if concepto == 'MAYORISTA': d['may'] = n(r[2])
        elif concepto == 'MATEU':
            d['edlp'], d['ruge'], d['acc'], d['ruge2da'], d['edlp2da'], d['edlpAnt'] = [n(v) for v in r[2:8]]
    wb.close()
    return out

cons = leer_consolidado(HIST)
for k, d in sorted(cons.items()):
    m = meses.get(k)
    retail = d['edlp']+d['ruge']+d['acc']+d['ruge2da']+d['edlp2da']+d['edlpAnt']
    ok = m and m['edlpUnidades']==d['edlp']+d['may'] and m['rugeUnidades']==d['ruge'] and m['retail']==retail and m['mayor']==d['may']
    print('consolidado', k, d, 'OK' if ok else '<-- NO CUADRA con meses')

out = {'meses':meses, 'consolidado':cons, 'rda':{k:rda[k] for k in sorted(rda)}}
js = "// Generado desde las liquidaciones reales de Juli (Excels 'Regalias Ruge SE y EDLP <mes> 2026'\n"
js += "// y 'DETALLE VENTA RUGE DEPORTES AMATEURS'). Regenerar con el script del chat si cambian.\n"
js += "// Marzo incluye el ajuste de +$363.780,63 (109 prendas EDLP 26 al 15% cuando correspondia 20%).\n"
js += "// Enero paga EDLP 26 al 20% y RUGE SE al 10% (escala 2025 vigente en el mes de transicion).\n"
js += "window.SIEMBRA_2026 = " + json.dumps(out, ensure_ascii=False, separators=(',',':')) + ";\n"
open(r"C:\Users\julia\Downloads\mateu-sports-portal\mateu-sports-portal\regalias\siembra-2026.js",'w',encoding='utf-8').write(js)
print('OK ->', len(js), 'bytes')
